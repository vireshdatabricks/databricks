"""Render a report snapshot package to HTML and Excel (reference/44 §10, decisions O3/O4).

Both formats are pure functions of the frozen package plus the review decisions recorded in the
application database (``review.export_inputs``). Item keys (``finding:CASE:field``, ``theme:ID``,
``scope:CASE``, ``section:ID``, ``query:ID``) match ``review.items_from_package``. Items without a
decision are UNVALIDATED and watermarked; ``validated_only`` keeps only VALIDATED and REVISED items,
and a REVISED item shows the reviewer's text.
"""
from __future__ import annotations

import html
import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

UNVALIDATED = "UNVALIDATED"
SHOWN_WHEN_VALIDATED_ONLY = {"VALIDATED", "REVISED"}
FINDING_FIELDS = [
    ("what_happened", "What happened"), ("what_failed", "What failed"), ("where_it_occurred", "Where it occurred"),
    ("how_detected", "How it was detected"), ("resolution", "Resolution"), ("prevention_opportunity", "Prevention opportunity"),
    ("likely_causes", "Likely causes"),
]


class Review:
    """Decision lookups shared by the HTML and Excel renderers."""

    def __init__(self, decisions: dict[str, dict[str, Any]] | None, validated_only: bool):
        self.decisions = decisions or {}
        self.validated_only = validated_only

    def status(self, key: str) -> str:
        return (self.decisions.get(key) or {}).get("status") or UNVALIDATED

    def shown(self, key: str) -> bool:
        return not self.validated_only or self.status(key) in SHOWN_WHEN_VALIDATED_ONLY

    def text(self, key: str, original: Any) -> Any:
        decision = self.decisions.get(key) or {}
        return decision.get("revised_text") if decision.get("status") == "REVISED" else original

    def comment(self, key: str) -> str:
        return (self.decisions.get(key) or {}).get("comment") or ""


def _finding_value(finding: dict[str, Any], field: str) -> str | None:
    if field == "likely_causes":
        return "; ".join(f"{c['cause']} ({c['cause_kind'].replace('_', ' ').lower()})" for c in finding.get("likely_causes") or []) or None
    return finding.get(field)


def _source(citation: dict[str, Any]) -> str:
    if citation.get("source_type") == "ATTACHMENT_CHUNK":
        return f"{citation['record_number']} · attachment {citation.get('attachment_name')} · {citation.get('page_or_sheet')}"
    if citation.get("source_type") == "CHILD_SUMMARY":
        return f"{citation['record_number']} · machine summary of {(citation.get('field_name') or '').replace('_', ' ')}"
    if citation.get("source_type") == "FINDING_FIELD":
        return f"{citation['record_number']} · finding field {citation.get('field_name')}"
    return f"{citation['record_number']} · field {citation.get('field_name')}"


def _flatten_citations(citations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Citations plus the attachment citations behind any task summary they point to."""
    rows = []
    for c in citations or []:
        rows.append({**c, "via": None})
        for nested in c.get("attachment_citations") or []:
            rows.append({**nested, "section": c.get("section"), "via": c.get("record_number")})
    return rows


def _in_scope_cases(package: dict[str, Any]) -> list[dict[str, Any]]:
    return [c for c in package["cases"].values() if (c.get("lens") or {}).get("in_scope") and c.get("finding")]


def build_section_data(package: dict[str, Any]) -> list[dict[str, Any]]:
    """Shared ordered section source for report rendering and the review document.

    The HTML renderer still owns markup; this builder makes the same frozen section order and
    content sources available to the review API without coupling it to generated HTML.
    """
    queries = {q.get("query_id"): q for q in package.get("queries") or []}
    result = []
    for section in package.get("sections") or []:
        blocks: list[dict[str, Any]] = []
        if section.get("text") and not section.get("generated"):
            blocks.append({"type": "text", "origin": "TEMPLATE", "text": section["text"], "generated": False})
        kind = section.get("kind")
        if kind == "themes":
            group_by = (package.get("promptbook") or {}).get("group_by") or []
            visible_themes = sorted(package.get("themes") or [], key=lambda t: (
                t.get("analysis_scope") != "CLIENT",
                *[(t.get("dominant_values") or {}).get(g, "") for g in group_by],
                -int(t.get("case_count") or 0)))
            for theme in visible_themes:
                cases = []
                for case_number in theme.get("case_numbers") or []:
                    case = package.get("cases", {}).get(case_number)
                    if not case or not (case.get("lens") or {}).get("in_scope") or not case.get("finding"):
                        continue
                    cases.append({"case_number": case_number,
                                  "item_keys": [f"finding:{case_number}:{field}" for field, _ in FINDING_FIELDS],
                                  "item_ids": [],
                                  "summary": {"undecided": 0, "validated": 0, "revised": 0, "rejected": 0}})
                blocks.append({"type": "theme_group", "theme_item_key": f"theme:{theme['theme_id']}",
                               "theme_item_id": None, "cases": cases})
        elif kind == "evidence_query":
            for ref in section.get("queries") or []:
                query = queries.get(ref.get("query_id")) or {}
                blocks.append({"type": "evidence_answer", "item_key": f"query:{ref.get('query_id')}",
                               "item_id": None, "query_id": ref.get("query_id"), "label": ref.get("label"),
                               "origin": "COMPUTED", "accepted": query.get("check_status") == "ACCEPTED",
                               "question": query.get("question"), "check_status": query.get("check_status"),
                               "check_detail": query.get("check_detail"), "generated_sql": query.get("generated_sql")})
            # Older package shape stores query.section_id without section.queries.
            if not section.get("queries"):
                for query in package.get("queries") or []:
                    if query.get("section_id") == section.get("id"):
                        blocks.append({"type": "evidence_answer", "item_key": f"query:{query.get('query_id')}",
                                       "item_id": None, "query_id": query.get("query_id"), "label": query.get("question"),
                                       "origin": "COMPUTED", "accepted": query.get("check_status") == "ACCEPTED",
                                       "question": query.get("question"), "check_status": query.get("check_status"),
                                       "check_detail": query.get("check_detail"), "generated_sql": query.get("generated_sql")})
        elif kind == "scope_appendix":
            for case in package.get("cases", {}).values():
                if case.get("lens"):
                    blocks.append({"type": "item", "item_key": f"scope:{case['case_number']}", "item_id": None})
        elif section.get("generated"):
            blocks.append({"type": "item", "item_key": f"section:{section['id']}", "item_id": None})
        if kind not in {"themes", "evidence_query", "scope_appendix"} and not section.get("generated"):
            if kind == "theme_table":
                blocks.append({"type": "facts", "origin": "COMPUTED", "rows": [
                    {"label": "Themes", "value": str((package.get("facts") or {}).get("theme_count", len(package.get("themes") or [])))},
                    {"label": "Cases in scope", "value": str((package.get("facts") or {}).get("in_scope_count", ""))}]})
        result.append({"template": section, "blocks": blocks})
    return result


# ---------------------------------------------------------------------------- HTML
_CSS = """
body{font-family:Segoe UI,Arial,sans-serif;color:#1f2933;max-width:1100px;margin:24px auto;padding:0 16px;line-height:1.45;background:#fff}
h1{font-size:24px;margin-bottom:4px} h2{border-bottom:2px solid #d9e2ec;padding-bottom:4px;margin-top:32px} h4{margin:16px 0 4px}
table{border-collapse:collapse;width:100%;margin:8px 0;font-size:14px} th,td{border:1px solid #d9e2ec;padding:6px 8px;text-align:left;vertical-align:top}
th{background:#f0f4f8} .banner{background:#fff4e5;border:1px solid #f0b429;padding:10px 14px;border-radius:6px}
.tag{display:inline-block;font-size:12px;background:#fff4e5;border:1px solid #f0b429;border-radius:10px;padding:0 8px;margin-left:4px}
.tag.fact{background:#e3f8ec;border-color:#3ebd93} .tag.model{background:#e6f0ff;border-color:#5a8dee}
.tag.VALIDATED{background:#e3f8ec;border-color:#3ebd93} .tag.REVISED{background:#e6f0ff;border-color:#5a8dee} .tag.REJECTED{background:#fde8e8;border-color:#e12d39}
.sub,.src{color:#52606d;font-size:13px} .warn{color:#b44d12} .cites{margin:6px 0 0 18px;font-size:13px} .nested{margin-left:18px}
.review{font-size:13px;width:150px} details{margin:8px 0;border:1px solid #d9e2ec;border-radius:6px;padding:6px 10px}
.original{color:#52606d;text-decoration:line-through;font-size:13px} pre{white-space:pre-wrap;font-size:12px}
@media print{details{break-inside:avoid}}
"""
_UNVALIDATED_TAG = '<span class="tag">Unvalidated &middot; likely contributing factor</span>'
_FACT_TAG = '<span class="tag fact">Computed fact</span>'
_MODEL_TAG = '<span class="tag model">Model classification</span>'


def _e(value: Any) -> str:
    return html.escape("" if value is None else str(value))


def _review_cell(review: Review, key: str) -> str:
    status = review.status(key)
    if status == UNVALIDATED:
        return "&#9744; Validated<br>&#9744; Revise<br>&#9744; Reject"
    decision = review.decisions[key]
    comment = f"<br><span class=\"sub\">{_e(decision.get('comment'))}</span>" if decision.get("comment") else ""
    return f'<span class="tag {status}">{status.title()}</span><br><span class="sub">{_e(decision.get("reviewer"))}</span>{comment}'


def _reviewed_text(review: Review, key: str, original: Any) -> str:
    if review.status(key) == "REVISED":
        return f"{_e(review.text(key, original))}<div class=\"original\">{_e(original)}</div>"
    return _e(original)


def _citation_list(citations: list[dict[str, Any]], section: str | None = None) -> str:
    items = []
    for c in citations or []:
        if section and c.get("section") != section:
            continue
        line = f'<li>&ldquo;{_e(c["excerpt"])}&rdquo; <span class="src">{_e(_source(c))}</span>'
        nested = c.get("attachment_citations") or []
        if nested:
            line += '<ul class="nested">' + "".join(
                f'<li>&ldquo;{_e(n["excerpt"])}&rdquo; <span class="src">{_e(_source(n))}</span></li>' for n in nested) + "</ul>"
        if c.get("note"):
            line += f' <span class="warn">({_e(c["note"])})</span>'
        items.append(line + "</li>")
    return f'<ul class="cites">{"".join(items)}</ul>' if items else '<p class="warn">No verified citation for this item.</p>'


def _case_block(case: dict[str, Any], lens_dims: list[dict[str, Any]], review: Review) -> str:
    finding, lens = case.get("finding") or {}, case.get("lens") or {}
    rows = []
    for field, label in FINDING_FIELDS:
        key = f"finding:{case['case_number']}:{field}"
        if not review.shown(key):
            continue
        value = _finding_value(finding, field)
        body = _reviewed_text(review, key, value) if value else '<span class="warn">Omitted: no verifiable citation.</span>'
        rows.append(f"<tr><th>{label}</th><td>{body}{_citation_list(finding.get('citations'), field)}</td>"
                    f'<td class="review">{_review_cell(review, key)}</td></tr>')
    if not rows:
        return ""
    lens_values = "; ".join(f"{d.get('label', d['name'])}: {', '.join((lens.get('values') or {}).get(d['name'], []))}" for d in lens_dims)
    path = " → ".join(finding.get("assignment_group_path") or [])
    return (f"<details open><summary><b>{_e(case['case_number'])}</b> · opened {_e((case.get('opened_at') or '')[:10])}"
            f" · closed {_e((case.get('closed_at') or '')[:10])} · {_e(case.get('tat_calendar_days'))} calendar days"
            f" · finding confidence {_e(finding.get('confidence'))}</summary>"
            f"<p class=\"sub\">Assignment path {_FACT_TAG}: {_e(path) or 'n/a'}<br>Classification {_MODEL_TAG}: {_e(lens_values)}</p>"
            + (f"<p class=\"sub\">Systems and artifacts named: {_e(', '.join(finding.get('systems_and_artifacts') or []))}</p>"
               if finding.get("systems_and_artifacts") else "")
            + f"<table><tr><th>Item</th><th>Candidate finding and evidence {_UNVALIDATED_TAG}</th><th>Reviewer</th></tr>{''.join(rows)}</table>"
            + (f"<p class=\"sub\">Evidence limitations: {_e(finding.get('evidence_limitations'))}</p>" if finding.get("evidence_limitations") else "")
            + "</details>")


def _visible_themes(package: dict[str, Any], review: Review) -> list[dict[str, Any]]:
    group_by = package["promptbook"]["group_by"]
    themes = [t for t in package["themes"] if review.shown(f"theme:{t['theme_id']}")]
    return sorted(themes, key=lambda t: (t["analysis_scope"] != "CLIENT", *[t["dominant_values"].get(g, "") for g in group_by], -t["case_count"]))


def _theme_table(package: dict[str, Any], review: Review) -> str:
    labels = {d["name"]: d.get("label", d["name"]) for d in package["promptbook"]["lens"]}
    group_by = package["promptbook"]["group_by"]
    head = "".join(f"<th>{_e(labels[g])}</th>" for g in group_by)
    rows = "".join(
        "<tr>" + "".join(f"<td>{_e(t['dominant_values'].get(g))}</td>" for g in group_by)
        + f"<td>{_e(t['primary_assignment_group'])}</td><td>{_e(t['theme_name_plain'])}</td>"
        f"<td>{t['case_count']} case{'s' if t['case_count'] != 1 else ''}, {round(100 * t['frequency_pct'])}%</td>"
        f"<td>{_e(t['key_systemic_action_plain'])}</td>"
        f"<td>{'Limited evidence' if t['limited_evidence'] else 'Meets minimum'} · {_e(t['confidence'])}</td>"
        f"<td>{_e(review.status('theme:' + t['theme_id']).title())}</td></tr>" for t in _visible_themes(package, review))
    return (f"<table><tr>{head}<th>Assignment group</th><th>Theme</th><th>Frequency {_FACT_TAG}</th>"
            f"<th>Key systemic action</th><th>Evidence</th><th>Review</th></tr>{rows}</table>")


def _themes(package: dict[str, Any], review: Review) -> str:
    out, lens_dims = [], package["promptbook"]["lens"]
    thresholds = package["promptbook"]["thresholds"]
    for i, t in enumerate(_visible_themes(package, review), 1):
        key = f"theme:{t['theme_id']}"
        recurrence = "".join(f"<tr><td>{_e(r['client_account'])}</td><td>{_e(r['report_month'][:7])}</td><td>{r['case_count']}</td>"
                             f"<td>{r['cumulative_case_count']}</td><td>{'Yes' if r['is_same_client_repeat'] else 'No'}</td></tr>" for r in t["recurrence"])
        flags = t.get("quality_flags") or []
        cases = [package["cases"][n] for n in t["case_numbers"] if n in package["cases"]]
        out.append(
            f"<section><h3>Theme {i}: {_e(t['theme_name_plain'])}</h3>"
            f"<p class=\"sub\">{_e(t['theme_name_technical'])} · {_e(t['analysis_scope'].replace('_', '-').lower())} · theme review: "
            f"{_review_cell(review, key)}</p>"
            f"<h4>Problem statement {_UNVALIDATED_TAG}</h4><p>{_reviewed_text(review, key, t['problem_statement_plain'])}</p>"
            f"<p class=\"sub\">{_e(t['problem_statement_technical'])}</p>"
            f"<h4>Frequency and timing {_FACT_TAG}</h4><p>{t['case_count']} of {t['in_scope_case_count']} in-scope cases "
            f"({round(100 * t['frequency_pct'])}%), out of {t['selected_case_count']} selected; first opened {_e((t['first_opened_at'] or '')[:10])}, "
            f"last opened {_e((t['last_opened_at'] or '')[:10])}."
            + (f" <b>Limited evidence:</b> fewer than {thresholds['min_theme_support']} supporting cases." if t["limited_evidence"] else "")
            + f"</p><h4>Attribution</h4><p>{_e(t.get('attribution_text'))} {_UNVALIDATED_TAG}</p>"
            "<p class=\"sub\">Dominant values (computed from case classifications): "
            + "; ".join(f"{_e(d.get('label', d['name']))} {_e(t['dominant_values'].get(d['name']))}" for d in lens_dims)
            + f" · primary assignment group {_e(t['primary_assignment_group'])}. Confidence {_e(t['confidence'])}: {_e(t['confidence_explanation'])}</p>"
            f"<h4>Supporting cases and evidence</h4>{''.join(_case_block(c, lens_dims, review) for c in cases)}"
            f"<h4>Systemic action {_UNVALIDATED_TAG}</h4><p>{_e(t['systemic_action_detail'])}</p>"
            f"<p><b>Why this fix works:</b> {_e(t['why_fix_works'])}</p>"
            f"<p><b>Implementing teams:</b> {_e(', '.join(t.get('implementing_teams') or []))} · "
            f"<b>Proposed timeline (reviewer sets final):</b> {_e(t['proposed_timeline'])}</p>"
            f"<h4>Recurrence {_FACT_TAG}</h4><table><tr><th>Client</th><th>Month</th><th>Cases</th><th>Cumulative</th>"
            f"<th>Repeat within {thresholds['same_client_repeat_window_days']} days</th></tr>{recurrence}</table>"
            + (f"<p class=\"warn\">Plain-language check: {_e('; '.join(flags))}</p>" if flags else "")
            + "</section>")
    return "".join(out) or '<p class="warn">No themes to show.</p>'


def _scope_appendix(package: dict[str, Any], review: Review) -> str:
    rows = "".join(
        f"<tr><td>{_e(c['case_number'])}</td><td>{_e(c['lens']['exclusion_reason'])}</td><td>{_e(c['lens']['scope_rationale'])}</td>"
        f"<td>{_e(c['lens']['confidence'])}</td><td class=\"review\">{_review_cell(review, 'scope:' + c['case_number'])}</td></tr>"
        for c in package["cases"].values() if c.get("lens") and not c["lens"]["in_scope"] and review.shown("scope:" + c["case_number"]))
    return (f"<table><tr><th>Case</th><th>Exclusion reason {_MODEL_TAG}</th><th>Rationale</th><th>Confidence</th><th>Reviewer</th></tr>"
            f"{rows or '<tr><td colspan=5>None</td></tr>'}</table>")


def _evidence_queries(package: dict[str, Any], section: dict[str, Any], review: Review) -> str:
    queries, out = {q["query_id"]: q for q in package.get("queries") or []}, []
    for ref in section.get("queries") or []:
        q, key = queries.get(ref["query_id"], {}), f"query:{ref['query_id']}"
        if q.get("check_status") == "ACCEPTED" and not review.shown(key):
            continue
        header = f"<h4>{_e(ref.get('label') or q.get('question'))} {_FACT_TAG if q.get('check_status') == 'ACCEPTED' else ''}</h4>"
        if q.get("check_status") != "ACCEPTED":
            if not review.validated_only:
                out.append(header + f'<p class="warn">Not used: {_e(q.get("check_status"))} — {_e(q.get("check_detail"))}.</p>'
                           + (f"<details><summary>SQL Genie produced</summary><pre>{_e(q.get('generated_sql'))}</pre></details>" if q.get("generated_sql") else ""))
            continue
        head = "".join(f"<th>{_e(c)}</th>" for c in q["result_columns"])
        rows = "".join("<tr>" + "".join(f"<td>{_e(v)}</td>" for v in row) + "</tr>" for row in q["rows"])
        out.append(header + f"<p class=\"sub\">Question: {_e(q['question'])}<br>Check: {_e(q['check_detail'])} · {q['result_row_count']} rows"
                   + (" (first rows shown)" if q["result_row_count"] > len(q["rows"]) else "") + "</p>"
                   f"<table><tr>{head}</tr>{rows}</table>"
                   f"<details><summary>SQL (generated by Genie, checked before use)</summary><pre>{_e(q['generated_sql'])}</pre></details>"
                   f"<p class=\"review\">{_review_cell(review, key)}</p>")
    return "".join(out) or '<p class="warn">No evidence answers to show.</p>'


def _limitations(package: dict[str, Any], review: Review) -> str:
    counts = ", ".join(f"{k.replace('_', ' ').lower()}: {v}" for k, v in sorted(package["attachment_counts"].items())) or "none linked"
    items = [*package["disclosures"], f"Attachments for these cases: {counts}. Only the most relevant attachment content is used for each task.",
             f"Themes with fewer than {package['promptbook']['thresholds']['min_theme_support']} supporting cases are marked limited evidence."]
    return "<ul>" + "".join(f"<li>{_e(i)}</li>" for i in items) + "</ul>"


def render_html(package: dict[str, Any], decisions: dict[str, dict[str, Any]] | None = None, validated_only: bool = False,
                report_status: str | None = None) -> str:
    review = Review(decisions, validated_only)
    run, book = package["run"], package["promptbook"]
    renderers = {"theme_table": _theme_table, "themes": _themes, "scope_appendix": _scope_appendix, "limitations": _limitations}
    body = []
    for section_source in build_section_data(package):
        section = section_source["template"]
        key = f"section:{section['id']}"
        if section.get("generated") and not review.shown(key):
            continue
        text = ""
        if section.get("text"):
            shown = _reviewed_text(review, key, section["text"]) if section.get("generated") else _e(section["text"])
            text = f"<p>{shown}</p>".replace("\n", "<br>")
        elif section.get("generated"):
            text = '<p class="warn">This section could not be generated.</p>'
        tag = (f"{_UNVALIDATED_TAG} <span class=\"review\">{_review_cell(review, key)}</span>" if section.get("generated") else "")
        rendered = (_evidence_queries(package, section, review) if section["kind"] == "evidence_query"
                    else renderers.get(section["kind"], lambda p, r: "")(package, review))
        body.append(f"<h2>{_e(section['heading'])}</h2>{tag}{text}{rendered}")
    status_line = f" · review status {_e(report_status)}" if report_status else ""
    banner = ("<b>Validated items only.</b> This export contains only items a reviewer validated or revised."
              if validated_only else
              "<b>Review draft.</b> Items marked unvalidated are machine-generated candidates until a reviewer validates them. Counts, "
              "dates, assignment paths, and recurrence are computed directly from ServiceNow data. Each finding lists the exact source "
              "text it relies on; a quoted excerpt is verified to exist in its source, but whether it supports the conclusion is for the "
              "reviewer to judge.")
    return f"""<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>{_e(book['title'])}</title><style>{_CSS}</style></head><body>
<h1>{_e(book['title'])}</h1>
<p class="sub">{_e(', '.join(run['clients']))} · analysis run {_e(run['analysis_run_id'])} · promptbook {_e(book['promptbook_id'])} v{_e(book['version'])}
· generated {_e(package['generated_at'][:16].replace('T', ' '))} UTC · model {_e(run['model_id'])}{status_line}</p>
<p class="banner">{banner}</p>
<p class="sub">How to read: {_FACT_TAG} · {_MODEL_TAG} · {_UNVALIDATED_TAG}</p>
{''.join(body)}
</body></html>"""


# ---------------------------------------------------------------------------- Excel
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="334E68")
_UNVALIDATED_FILL = PatternFill("solid", fgColor="FFF4E5")
_REVIEW_HEADERS = ["Validation status", "Reviewer decision text", "Reviewer comment"]


def _sheet(workbook: Workbook, title: str, headers: list[str], rows: list[list[Any]], widths: dict[int, int] | None = None):
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font, cell.fill = _HEADER_FONT, _HEADER_FILL
    status_col = headers.index("Validation status") + 1 if "Validation status" in headers else None
    for row in rows:
        sheet.append(["; ".join(map(str, v)) if isinstance(v, list) else v for v in row])
        if status_col and sheet.cell(sheet.max_row, status_col).value == UNVALIDATED:
            for cell in sheet[sheet.max_row]:
                cell.fill = _UNVALIDATED_FILL
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = (widths or {}).get(index, min(max(len(header) + 4, 14), 60))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    return sheet


def _review_columns(review: Review, key: str) -> list[Any]:
    decision = review.decisions.get(key) or {}
    return [review.status(key), decision.get("revised_text") if decision.get("status") == "REVISED" else "", review.comment(key)]


def render_excel(package: dict[str, Any], decisions: dict[str, dict[str, Any]] | None = None, validated_only: bool = False,
                 report_status: str | None = None) -> bytes:
    review = Review(decisions, validated_only)
    run, book, facts = package["run"], package["promptbook"], package["facts"]
    lens_dims = book["lens"]
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append([book["title"]])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append(["VALIDATED ITEMS ONLY" if validated_only else
                    "REVIEW DRAFT: rows marked UNVALIDATED are machine-generated candidates until a reviewer validates them."])
    summary["A2"].fill = _UNVALIDATED_FILL
    for label, value in [("Client(s)", ", ".join(run["clients"])), ("Analysis run", run["analysis_run_id"]),
                         ("Promptbook", f"{book['promptbook_id']} v{book['version']}"), ("Requested window", f"{run['date_from']} to {run['date_to']}"),
                         ("Cases selected", facts["selected_count"]), ("Cases in scope (model classification)", facts["in_scope_count"]),
                         ("Themes", facts["theme_count"]), ("Data through", facts["data_through"]), ("Model", run["model_id"]),
                         ("Review status", report_status or "not imported for review"),
                         ("Export filter", "validated items only" if validated_only else "all items, unvalidated watermarked")]:
        summary.append([label, value])
    summary.append([])
    for section in package["sections"]:
        key = f"section:{section['id']}"
        if section.get("text") and (not section.get("generated") or review.shown(key)):
            summary.append([section["heading"], review.text(key, section["text"]) if section.get("generated") else section["text"]])
    summary.append([])
    summary.append(["Limitations"])
    for item in package["disclosures"]:
        summary.append(["", item])
    summary.column_dimensions["A"].width, summary.column_dimensions["B"].width = 38, 120
    for row in summary.iter_rows(min_row=3):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    theme_rows = []
    for t in package["themes"]:
        key = f"theme:{t['theme_id']}"
        if review.shown(key):
            theme_rows.append([t["theme_name_plain"], *[t["dominant_values"].get(d["name"]) for d in lens_dims], t["primary_assignment_group"],
                               t["case_count"], t["in_scope_case_count"], round(t["frequency_pct"], 4), (t["first_opened_at"] or "")[:10],
                               (t["last_opened_at"] or "")[:10], "Yes" if t["limited_evidence"] else "No", t["confidence"],
                               t.get("attribution_text"), t["problem_statement_plain"], t["key_systemic_action_plain"], t["systemic_action_detail"],
                               t["why_fix_works"], t.get("implementing_teams") or [], t["proposed_timeline"], t["case_numbers"],
                               t.get("quality_flags") or [], *_review_columns(review, key)])
    _sheet(workbook, "Themes", ["Theme", *[d.get("label", d["name"]) for d in lens_dims], "Primary assignment group", "Cases", "In-scope cases",
                                "Share of in-scope", "First opened", "Last opened", "Limited evidence", "Confidence", "Attribution (computed)",
                                "Problem statement", "Key systemic action", "Systemic action detail", "Why it works", "Implementing teams",
                                "Proposed timeline", "Supporting cases", "Plain-language flags", *_REVIEW_HEADERS], theme_rows)

    finding_rows, citation_rows = [], []
    for case in _in_scope_cases(package):
        finding = case["finding"]
        for field, label in FINDING_FIELDS:
            key = f"finding:{case['case_number']}:{field}"
            if review.shown(key):
                cited = [c for c in _flatten_citations(finding.get("citations")) if c.get("section") == field]
                finding_rows.append([case["case_number"], case["client_account"], label,
                                     _finding_value(finding, field) or "(omitted: no verifiable citation)", len(cited),
                                     finding.get("confidence"), *_review_columns(review, key)])
        for c in _flatten_citations(finding.get("citations")):
            if review.shown(f"finding:{case['case_number']}:{c.get('section')}"):
                citation_rows.append([case["case_number"], c.get("section"), c.get("source_type"), c.get("record_number"), c.get("field_name"),
                                      c.get("attachment_name"), c.get("page_or_sheet"), c.get("via"), c.get("excerpt")])
    _sheet(workbook, "Findings", ["Case", "Client", "Item", "Candidate finding", "Citations", "Finding confidence", *_REVIEW_HEADERS],
           finding_rows, {4: 80})
    _sheet(workbook, "Citations", ["Case", "Supports", "Source type", "Record", "Field", "Attachment", "Sheet / location",
                                   "Via task summary", "Verbatim excerpt"], citation_rows, {9: 80})

    case_rows = []
    for case in package["cases"].values():
        finding, lens = case.get("finding") or {}, case.get("lens") or {}
        key = f"scope:{case['case_number']}"
        if lens and review.shown(key):
            case_rows.append([case["case_number"], case["client_account"], (case.get("opened_at") or "")[:10], (case.get("closed_at") or "")[:10],
                              case.get("tat_calendar_days"), case.get("category"), case.get("subtype"), finding.get("assignment_group_path") or [],
                              "Yes" if lens.get("in_scope") else "No", lens.get("exclusion_reason"), lens.get("scope_rationale"),
                              *[(lens.get("values") or {}).get(d["name"], []) for d in lens_dims], *_review_columns(review, key)])
    _sheet(workbook, "Cases and scope", ["Case", "Client", "Opened", "Closed", "Calendar days", "Category (ServiceNow)", "Subtype (ServiceNow)",
                                         "Assignment path", "In scope (model)", "Exclusion reason", "Scope rationale",
                                         *[f"{d.get('label', d['name'])} (model)" for d in lens_dims], *_REVIEW_HEADERS], case_rows)
    _sheet(workbook, "Recurrence", ["Theme", "Client", "Month", "Cases", "Cumulative", "Same-client repeat", "Clients with theme", "Cross-client"],
           [[t["theme_name_plain"], r["client_account"], r["report_month"][:7], r["case_count"], r["cumulative_case_count"],
             "Yes" if r["is_same_client_repeat"] else "No", r["clients_with_theme_in_window"], "Yes" if r["is_cross_client"] else "No"]
            for t in package["themes"] if review.shown(f"theme:{t['theme_id']}") for r in t["recurrence"]])

    query_rows, result_rows = [], []
    for q in package.get("queries") or []:
        key = f"query:{q['query_id']}"
        accepted = q.get("check_status") == "ACCEPTED"
        if (accepted and review.shown(key)) or (not accepted and not validated_only):
            query_rows.append([q["section_id"], q["question"], q["check_status"], q.get("check_detail"), q.get("result_row_count"),
                               q.get("generated_sql"), q["query_id"], *(_review_columns(review, key) if accepted else ["NOT USED", "", ""])])
        if accepted and review.shown(key):
            result_rows.extend([q["question"], [f"{c}={v}" for c, v in zip(q["result_columns"], row)]] for row in q["rows"])
    _sheet(workbook, "Queries", ["Section", "Question", "Check", "Check detail", "Rows", "SQL (Genie, checked)", "Query id", *_REVIEW_HEADERS],
           query_rows, {2: 60, 6: 80})
    _sheet(workbook, "Query results", ["Question", "Values"], result_rows, {1: 60, 2: 80})
    processing = package["processing"]
    _sheet(workbook, "Processing", ["Item", "Value"],
           [["Model calls", processing.get("calls")], ["Prompt tokens", processing.get("prompt_tokens")],
            ["Completion tokens", processing.get("completion_tokens")], ["Failures", processing.get("failures") or []],
            ["Summary prompt version", processing.get("summary_prompt_version")],
            ["Finding prompt version", processing.get("finding_prompt_version")], ["Lens hash", processing.get("lens_hash")],
            ["Promptbook resolved hash", book["resolved_hash"]], ["Snapshot week", run["as_of_extract_week"]]], {2: 80})
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
